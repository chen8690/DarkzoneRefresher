import openpyxl, io, sys, zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from darkzone.rules import default_rules_path

path = default_rules_path()
out = io.StringIO()

wb = openpyxl.load_workbook(path, data_only=True)
out.write("SHEETS: %r\n" % (wb.sheetnames,))
for ws in wb.worksheets:
    out.write("=" * 70 + "\n")
    out.write("SHEET: %s dims=%s max_row=%s max_col=%s\n" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
    out.write("=" * 70 + "\n")
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out.write("%s = %r\n" % (c.coordinate, c.value))
        out.write("-- row %d --\n" % row[0].row)
    out.write("MERGED: %r\n" % ([str(r) for r in ws.merged_cells.ranges],))
    # data validation
    try:
        dvs = ws.data_validations.dataValidation
        out.write("DATA_VALIDATIONS count=%d\n" % len(dvs))
        for dv in dvs:
            out.write("  type=%r formula1=%r formula2=%r sqref=%r allow_blank=%r\n" % (dv.type, dv.formula1, dv.formula2, dv.sqref, dv.allow_blank))
    except Exception as e:
        out.write("DV error: %r\n" % (e,))
    out.write("COLS: %r\n" % ([ (k, v.width) for k,v in ws.column_dimensions.items() ],))

# raw xml parts
out.write("\n" + "=" * 70 + "\nRAW PARTS\n" + "=" * 70 + "\n")
with zipfile.ZipFile(path) as z:
    for n in z.namelist():
        out.write("PART: %s\n" % n)
    for n in z.namelist():
        if n.endswith(".xml"):
            data = z.read(n).decode("utf-8", "replace")
            out.write("\n---- %s (%d chars) ----\n" % (n, len(data)))
            out.write(data[:20000])
            out.write("\n")

with open(Path(__file__).resolve().parent / "xlsx_dump.txt", "w", encoding="utf-8") as f:
    f.write(out.getvalue())
print("done", len(out.getvalue()))
